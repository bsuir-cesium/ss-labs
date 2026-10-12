"""Create the editable report from the verified example, manual counts and real UI capture."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
D=Document();sec=D.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(2.2);sec.right_margin=Cm(1.8);sec.top_margin=Cm(1.5);sec.bottom_margin=Cm(1.5)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=D.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(12)
D.styles['Normal'].paragraph_format.space_after=Pt(7)
D.styles['Normal'].paragraph_format.line_spacing=1.1
D.styles['Heading 1'].font.size=Pt(16)
D.styles['Heading 1'].paragraph_format.space_after=Pt(12)
D.styles['Heading 2'].font.size=Pt(13)
D.core_properties.title='Лабораторная 2. Метрика Джилба для TypeScript'
D.core_properties.subject='Метрики сложности потока управления'
D.core_properties.author='Пузик Владислав; Хренков Даниил'

def p(text='',bold=False):
    par=D.add_paragraph();r=par.add_run(text);r.bold=bold;return par

def title(text):D.add_heading(text,level=1)
def page():
    par=D.add_paragraph()
    par.paragraph_format.page_break_before=True
    par.paragraph_format.space_after=Pt(0)
    par.paragraph_format.line_spacing=Pt(1)
    par.add_run().font.size=Pt(1)
def table(headers, rows, widths):
    t=D.add_table(rows=1, cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);tcPr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                tag=OxmlElement('w:'+side);tag.set(qn('w:w'),'75');tag.set(qn('w:type'),'dxa');margins.append(tag)
            tcPr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                tag=OxmlElement('w:'+edge);tag.set(qn('w:val'),'single');tag.set(qn('w:sz'),'4');tag.set(qn('w:color'),'D9D9D9');borders.append(tag)
            tcPr.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF3');tcPr.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(1);par.paragraph_format.line_spacing=1
                for run in par.runs:run.font.size=Pt(10.5);run.bold=ri==0
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    D.add_paragraph().paragraph_format.space_after=Pt(0)

p('Дисциплина: метрология программного обеспечения').alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(5):p()
q=p('ЛАБОРАТОРНАЯ РАБОТА № 2',True);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.runs[0].font.size=Pt(20)
q=p('Метрики сложности потока управления');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.runs[0].font.size=Pt(16)
p('Анализ программы на TypeScript по метрике Джилба').alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(4):p()
p('Выполнили: Пузик Владислав')
p('Хренков Даниил')
p('Группа: 451001')
for _ in range(5):p()
p('2026').alignment=WD_ALIGN_PARAGRAPH.CENTER
page()
title('1. Цель и теоретические сведения')
p('Цель: подготовить программу со всеми видами циклов TypeScript и ветвлениями, построить подробную схему, вручную рассчитать метрику Джилба и проверить результат программным анализатором.')
D.add_heading('Метрика Джилба',level=2)
p('Абсолютная сложность CL — количество ветвлений исходной структуры. Учитываются if, каждый цикл и каждый case. Относительная сложность cl = CL / N, где N — число операторов по принятому профилю. Максимальная вложенность CLI — наибольший уровень ветвления; внешний уровень равен нулю. При CL > 0 выполняется CL > CLI; без ветвлений оба равны нулю.')
p('Для расчёта вложенности switch представляется цепочкой вложенных if. Селектор и default не добавляют условий. Для k case глубины сравнений равны d, d+1, …, d+k−1. Тело j-го case имеет глубину d+j (j с 1). Поэтому отдельный switch с четырьмя case даёт CL=4 и CLI=3. Else и завершающий while в do-while не считаются отдельными ветвлениями.')
p('В N входят самостоятельные выражения, объявления переменных, управляющие операторы и переходы. Блоки, метки, типы и объявления функций не учитываются; тела функций анализируются. Заголовок for входит в один оператор цикла. При N=0 относительная сложность не определена.')
D.add_heading('Метрика Маккейба',level=2)
p('Цикломатическая сложность характеризует число линейно независимых путей графа потока управления: V(G)=E−V+2 для одной функции с нормализованными входом и выходом; для P таких компонент — E−V+2P. E — число дуг, V — число вершин. При бинарных решениях V(G)=D+1. Многовыходный узел добавляет число исходов минус один.')
p('В связном планарном изображении сложность также равна числу областей, включая внешнюю. Это размер базиса путей, а не число всех возможных трасс. Циклы допускают неограниченное число трасс; некоторые пути могут быть невыполнимы из-за зависимых условий. Метрика помогает планировать структурное тестирование, но не гарантирует качество программы.')
p('Для схемы данного примера D=12, поэтому V(G)=13. Это теоретический расчёт: приложение реализует Джилба. При раскрытии короткого замыкания логических выражений граф и D могут отличаться. Одинаковая V(G) не означает одинаковую вложенность. Джилб и Маккейб не измеряют полноту требований, сложность данных или предметной области.')
page()
title('2. Исходный код анализируемой программы')
p('Файл example.ts. Программа использует пять видов циклов: for, while, do-while, for-of и for-in. Числа слева — номера строк, они не входят в код.')
code=D.add_paragraph();code.paragraph_format.space_after=Pt(0);code.paragraph_format.line_spacing=Pt(11.1)
for i,line in enumerate((ROOT.parent/'example.ts').read_text().splitlines(),1):
    r=code.add_run(f'{i:02}  {line}\n');r.font.name='Courier New';r.font.size=Pt(9)
p('При отдельном запуске выводится total = 20. Метрики статические: количество исполненных итераций и достижимость отдельных case при выбранных константах не меняют структуру исходного текста.')
page()
title('3. Ручной расчёт')
rows=[(4,'for',0),(5,'if (i > 0)',1),(7,'case 1',2),(8,'if (total === 0)',3),(14,'case 2',3),(17,'case 3',4),(20,'case 4',5),(31,'while',0),(36,'do-while',0),(40,'for-of',0),(43,'for-in',0),(46,'if (total < 0)',0)]
table(['Строка','Ветвление','Уровень'],rows,[2,11,3.5])
p('CL = 5 циклов + 3 if + 4 case = 12. Максимум в таблице: CLI = 5. Вложенный if первого case имеет уровень 3; добавлять его глубину ко всему switch нельзя.')
table(['Категория','Строки исходника','Количество'],[
('Объявления','1, 2, 3, 30, 35',5),('Выражения','9, 11, 15, 18, 21, 24, 27, 32, 33, 37, 38, 41, 44, 47, 49',15),
('if','5, 8, 46',3),('Циклы','4, 31, 36, 40, 43',5),('switch','6',1),('break','13, 16, 19, 22',4)], [3.3,10.5,2.7])
p('N = 5 + 15 + 3 + 5 + 1 + 4 = 33.',True)
p('cl = 12 / 33 = 4 / 11 ≈ 0,363636; CL = 12; CLI = 5.',True)
p('Проверка: 12 > 5. Инициализация let i=0 внутри for не является самостоятельным Statement в TS. Скрытые r, t и keys на раскрытой схеме в N не входят. Схема следует на четырёх листах; одинаковые буквы соединителей обозначают продолжение потока.')
for i in range(1,5):
    page()
    D.add_picture(str(ROOT/f'flowchart-{i}.png'),width=Cm(16.4))
page()
title('4. Программный расчёт и проверка')
p('Графический анализатор разбирает example.ts через TypeScript Compiler API 5.9.3 в браузере. Синтаксическое дерево обходится рекурсивно; у каждого case отдельно вычисляется глубина тела. Исходный код не исполняется, импорты не загружаются.')
D.add_picture(str(ROOT/'analyzer-screenshot.jpg'),width=Cm(16.5))
p('Рисунок — фактический результат работы анализатора на example.ts.',False)
table(['Показатель','Вручную','Программой'],[('CL',12,12),('N',33,33),('cl','12/33 ≈ 0,363636','0,363636'),('CLI',5,5)],[8.5,4,4])
p('Проверены: if в первом и последнем case, вложенный switch, default в середине, fallthrough, else без фигурных скобок, все циклы, ASI, комментарии, строки, регулярные выражения, шаблонные строки и независимая вложенность функций. Синтаксические ошибки и неподдерживаемые конструкции показываются вместо результатов.')

page()
title('5. Вывод')
p('Ограничения профиля: не поддерживаются классы, методы объектов, try/catch/finally, генераторы, метки управления, namespace/enum и tagged templates. Проверяется синтаксис, но не типы/существование имён. Приложение не является универсальным анализатором всех расширений TypeScript.')
p('Созданы анализируемая программа на TypeScript, графический анализатор метрики Джилба и подробная схема. Ручной и программный расчёты совпали: CL=12, N=33, cl=12/33, CLI=5. Самая глубокая точка — сравнение четвёртого case. Вложенный if первого case не увеличивает этот максимум.')
D.add_heading('Контрольные вопросы',level=2)
for q in [
'Почему do-while даёт одно ветвление? Потому что это один оператор цикла с проверкой после тела.',
'Что изменит перенос if из первого case в четвёртый? CL и N останутся прежними, CLI увеличится с 5 до 6.',
'Почему default не добавляет CL? В нём не проверяется новое условие; собственные if/циклы в его теле учитываются.',
'Почему N не равен числу точек с запятой? Заголовки циклов, ASI и составные операторы определяются грамматикой.',
'Чем for-in отличается от for-of? Первый перечисляет строковые ключи, второй получает значения итератора.',
'Почему Маккейб не равен числу всех путей? Он задаёт размер независимого базиса, а циклы допускают повторение трасс.'
]:p(q)
D.save(ROOT/'report.docx')
print(ROOT/'report.docx')
